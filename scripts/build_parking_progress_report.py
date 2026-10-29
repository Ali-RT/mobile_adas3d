"""Build the dated Word and HTML report addendum from reviewed result facts.

Run with the bundled document-authoring Python runtime (python-docx required).
Earlier reports are never modified.
"""
from pathlib import Path
import html
import json

ROOT = Path(__file__).resolve().parents[1]
STEM = "MobileADAS3D_near_field_update_2026-10-29"


def content(facts):
    return [
        ("Progress and current decision", [
            "We have a trained shared parking model, a complete near-field data census, and a frozen benchmark for measuring the next improvement. We selected the actual-camera adapter at epoch 4 as V0. Its weights will remain unchanged while we measure the new development baseline. No V0 or V1 accuracy score has yet been produced on this new benchmark.",
            "The use case is a camera mounted on a slowly moving vehicle, providing Vehicle and Pedestrian labels and distance estimates for ultrasonic-sensor validation. The critical distance band is 2 to 6 metres. We also retain 0 to 2 metre coverage so the closest range is not forgotten. Projected 3D cuboids are optional.",
            "This update continues the October 28 report. The earlier A2 and A2 Lite experiments remain historical baselines; this work focuses on a shared Urban nano image backbone with a trained depth decoder and camera-aware adapter. We have not replaced the iPhone app with the selected V0 checkpoint."
        ], [(["Work item", "Status", "Evidence"], [
            ["Near-range depth training", "Completed", "Epoch 7 selected from 10 epochs per arm"],
            ["Joint camera-aware pilot", "Completed", "Epoch 0 retained; no selected improvement"],
            ["Adapter-only full-pool pilot", "Completed", "Actual-camera epoch 4 selected"],
            ["Full nuScenes coverage census", "Completed", "34,149 keyframes across 850 scenes"],
            ["Frozen V0/V1 benchmark", "Prepared", "512 dev and 512 held-out images"],
            ["V0 score on the new benchmark", "Next", "No model inference yet"],
        ])]),
        ("What the model experiments showed", [
            "The shared model processes an image once and reuses three feature scales for detection, projected cuboid corners and depth. The detection body and its BatchNorm statistics stay frozen during these depth pilots. The depth decoder was trained using sparse LiDAR ground truth, not a teacher model. A small camera adapter then uses the camera calibration to adjust depth predictions.",
            "The parking-balanced depth recipe reduced the old development error from 2.142 m for its broad-loss control to 0.760 m. Training the camera-aware decoder and adapter together did not improve the selected checkpoint. Training only the 1,553 adapter parameters gave a selected error of 0.716 m, about 5.8% lower than the 0.760 m initialization. The original 57,441 depth parameters and detector stayed frozen.",
            "These numbers use the earlier 256-image development set and an image-macro four-cell metric. They are not the new benchmark score and cannot be directly compared with it. The adapter did not pass every earlier research check: Vehicle at 2 to 4 metres remained worse than a constant 4 metre diagnostic. This is a useful improvement, not deployment qualification."
        ], [(["Experiment", "Budget per arm", "Selected epoch", "Old dev MAE"], [
            ["Broad-loss control", "10 epochs / 1,920 updates", "4", "2.142 m"],
            ["Parking-balanced depth", "10 epochs / 1,920 updates", "7", "0.760 m"],
            ["Joint camera pilot", "5 epochs / 960 updates", "0", "0.760 m"],
            ["Adapter fixed-camera control", "5 epochs / 1,920 updates", "3", "0.759 m"],
            ["Adapter actual camera  V0", "5 epochs / 1,920 updates", "4", "0.716 m"],
        ])]),
        ("What data we now have", [
            "We acquired the nuScenes v1.0 trainval metadata, LiDAR semantic labels and all ten keyframe-only sensor parts in Google Drive. The 12 compressed archives total 45.53 GB. The complete census checked 204,894 image headers and projected 204,889 camera views. Five views were excluded because the camera and LiDAR timestamps differed by more than 50 milliseconds.",
            "Training, development and evaluation roles are kept apart by recording log. The assigned pools contain 477 training scenes, 150 development scenes and 113 evaluation scenes. Another 110 scenes from Mini recording logs are excluded. Three previously inspected evaluation logs are excluded from the new final benchmark, leaving 98 scenes in 11 uninspected logs. These are our frozen custom roles, not a claim to the standard official nuScenes benchmark split.",
            "Counts below are object observations at individual keyframes, supported by at least five conservative object-owned LiDAR points in a camera and range band. A moving sequence can contain many observations of the same tracked object. Unique tracks show the number of distinct annotated instances. Extended objects can occupy several bands, so rows must not be added together."
        ], [(["Training band", "Vehicle observations", "Vehicle tracks", "Pedestrian observations", "Pedestrian tracks"], [
            ["0 to 1 m", "1", "1", "24", "5"],
            ["0 to 2 m", "2,290", "578", "110", "43"],
            ["2 to 4 m", "10,930", "2,043", "1,049", "347"],
            ["4 to 6 m", "19,186", "3,142", "2,742", "843"],
            ["0 to 6 m", "20,366", "3,154", "3,573", "863"],
        ])]),
        ("What the coverage means", [
            "Vehicle observations outnumber Pedestrian observations by about 5.7 to 1 across 0 to 6 metres and 20.8 to 1 below 2 metres. The larger pool gives us substantially more training material in the critical 2 to 6 metre range, but it does not remove the closest-range pedestrian gap.",
            "The 0 to 1 metre region is too sparse for a credible qualification claim, especially for vehicles. The new development benchmark contains no supported Vehicle observations in that band. Pedestrian below 2 metres is also small: 12 observations from eight tracks and three logs in development. We will report this limitation rather than convert a few examples into a pass or fail.",
            "The reference is camera-axis depth in metres at visible LiDAR surface points. It is not distance to an object centre, radial range from the lens, or the nearest surface seen by an ultrasonic sensor. Sensor validation will still need camera and sensor alignment, timing checks and a defined target surface. A lower depth error alone does not establish ultrasonic equivalence."
        ], [(["Benchmark cell", "Dev observations / tracks / logs", "Held-out observations / tracks / logs"], [
            ["Vehicle 2 to 4 m", "262 / 131 / 12", "232 / 121 / 11"],
            ["Vehicle 4 to 6 m", "314 / 191 / 12", "308 / 171 / 11"],
            ["Pedestrian 2 to 4 m", "132 / 63 / 12", "140 / 58 / 9"],
            ["Pedestrian 4 to 6 m", "231 / 110 / 12", "224 / 106 / 11"],
            ["Vehicle 0 to 2 m", "120 / 55 / 12", "102 / 53 / 9"],
            ["Pedestrian 0 to 2 m", "12 / 8 / 3", "61 / 13 / 6"],
        ])]),
        ("How we will compare V0 and V1", [
            "The frozen package contains 512 development images and 512 held-out images, covering 1,009 unique keyframes. Its 39 immutable image and sparse-ground-truth shards total about 156 MB. Selection used ground-truth coverage only, retained rare close-range views, and did not consult model predictions. This deliberately enriched sample is not an estimate of natural traffic prevalence.",
            "Our primary metric is mean absolute depth error in metres, averaged equally across Vehicle and Pedestrian in the 2 to 4 and 4 to 6 metre bands. Within each cell, cameras are averaged within an object observation, observations within a tracked object, tracks within a recording log, and then logs equally. This reduces domination by dense point clouds, repeated frames or long logs.",
            "We also report pointwise mean, median and 90th-percentile absolute error, prediction coverage and the 0 to 2 metre breakdown. The 0 to 1 metre result is exploratory. Invalid or missing predictions remain visible in the denominator; later V0/V1 comparisons will also use exactly the same valid points. We will not fit an individual scale to each image, clamp predictions into the ground-truth band, or interpolate missing ground truth.",
            "This is a depth comparison on ground-truth-labelled object surfaces. It does not measure detection mAP, missed objects or cuboid quality. There is no plus-or-minus 10 centimetre accuracy rule and no product pass or fail threshold in this benchmark. We must set any eventual acceptance rule from the sensor-validation requirement, not from whichever model result looks best."
        ], []),
        ("Next actions and reporting boundaries", [
            "First, freeze V0 as the unchanged actual-camera epoch-4 checkpoint. Register its checkpoint hash, architecture, preprocessing, calibration mapping, depth sampling and runtime. Run the 512 development images only, save resumable predictions, and publish errors alongside prediction coverage. This is inference only: no optimizer updates, BatchNorm changes or checkpoint selection.",
            "Second, inspect the development errors by class, range, track and recording log. Use those findings to write a prospective V1 training recipe using training logs only. Keep the architecture and scoring identities fixed for a weight-only comparison. Extra epochs, class sampling or new losses should be justified by this baseline, not assumed to help.",
            "Third, compare the chosen V1 with V0 on development using paired points. Freeze the final candidate before opening the held-out benchmark. Evaluate on the iPhone separately for latency, heat and sustained camera operation. No V1 training, held-out model scoring or phone replacement has started in this update.",
            "For next week, the completed deliverables are the trained shared-model pilots, the full 0 to 6 metre census, explicit coverage gaps and a reproducible benchmark package. The next evidence we need is the new V0 development score. We should not claim a new benchmark improvement or a production-ready distance model before those measurements exist."
        ], []),
    ]


def build():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    facts = json.loads((ROOT / "artifacts/parking_v0_benchmark_progress_20261029.json").read_text())
    sections = content(facts)
    out = ROOT / "reports/MobileADAS3D"; out.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.7)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = doc.styles[name]; style.font.name = "Arial"; style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(9)
    doc.add_paragraph("MobileADAS3D near field progress update", "Title")
    doc.add_paragraph("29 October 2026", "Subtitle")
    body = ['<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
        '<title>MobileADAS3D near field progress update</title><style>body{max-width:1060px;margin:40px auto;padding:0 24px;font:17px/1.65 system-ui;color:#171c25;background:#fff}h1{font-size:34px;line-height:1.2}h2{margin-top:2em;line-height:1.25}nav{display:flex;gap:14px;flex-wrap:wrap}a{color:#17539a}table{border-collapse:collapse;width:100%;margin:20px 0 28px;font-size:15px}th,td{border:1px solid #d9d9d9;padding:11px;text-align:left}th{background:#e8eef4}td:not(:first-child){text-align:center}p{max-width:90ch}.meta{color:#526071}details{margin:20px 0}code{overflow-wrap:anywhere;font-size:13px}@media print{nav,button{display:none}body{margin:0;font-size:11pt}h2{break-after:avoid}tr{break-inside:avoid}}</style>',
        '<h1>MobileADAS3D near field progress update</h1><p class="meta">29 October 2026 · Addendum to the October 28 report</p><button onclick="window.print()">Print or save as PDF</button><nav>']
    body += [f'<a href="#s{i}">{html.escape(title)}</a>' for i, (title, _, _) in enumerate(sections)] + ['</nav>']
    for i, (title, paragraphs, tables) in enumerate(sections):
        if i:
            doc.add_page_break()
        doc.add_heading(title, 1); body.append(f'<section id="s{i}"><h2>{html.escape(title)}</h2>')
        for text in paragraphs:
            doc.add_paragraph(text); body.append(f'<p>{html.escape(text)}</p>')
        for headers, rows in tables:
            table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
            widths = ([1.9, 1.35, 3.85] if i == 0 else [2.4, 2, 1.1, 1.6] if i == 1 else
                      [1.1, 1.5, 1.2, 1.8, 1.5] if i == 2 else [2.3, 2.4, 2.4])
            for col, width in zip(table.columns, widths): col.width = Inches(width)
            for values, row in [(headers, table.rows[0])] + [(values, table.add_row()) for values in rows]:
                for j, (value, cell) in enumerate(zip(values, row.cells)):
                    cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ["top", "left", "bottom", "right"]:
                        el = OxmlElement("w:"+side); el.set(qn("w:w"), "90"); el.set(qn("w:type"), "dxa"); margins.append(el)
                    props.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs: run.font.size = Pt(9.5)
            borders = OxmlElement("w:tblBorders")
            for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                el = OxmlElement("w:"+side); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); borders.append(el)
            table._tbl.tblPr.append(borders)
            repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell in table.rows[0].cells:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E8EEF4"); cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs: run.bold = True
            body.append('<table><thead><tr>'+''.join(f'<th>{html.escape(h)}</th>' for h in headers)+'</tr></thead><tbody>')
            for row in rows: body.append('<tr>'+''.join(f'<td>{html.escape(v)}</td>' for v in row)+'</tr>')
            body.append('</tbody></table>')
        body.append('</section>')
    doc.core_properties.title = "MobileADAS3D near field progress update"
    doc.core_properties.subject = "Training evidence and frozen V0 V1 benchmark readiness"
    doc.core_properties.author = "MobileADAS3D project"
    doc.add_paragraph("Evidence retained", "Heading 2")
    doc.add_paragraph("Reviewed result bundles: parking_shared_near_results.zip, parking_shared_camera_results.zip, parking_shared_adapter_results.zip, nuscenes_0_6m_coverage_results.zip and parking_v0_v1_benchmark_results.zip. Download completion is recorded in download_manifest_keyframes_all.json. Full hashes and frozen identities accompany the HTML report and machine-readable evidence record.")
    doc.save(out / (STEM + ".docx"))
    body.append('<h2>Evidence and reproducibility</h2><p>Source result bundles are retained separately from this report. The frozen data and metric identities allow the next model result to be traced to the same benchmark.</p><details><summary>Show source hashes and frozen identities</summary><pre><code>'+html.escape(json.dumps(facts, indent=2))+'</code></pre></details></html>')
    (out / (STEM + ".html")).write_text('\n'.join(body))
    print(out / (STEM + ".docx")); print(out / (STEM + ".html"))


if __name__ == "__main__":
    build()
